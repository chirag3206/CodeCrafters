"""
PeoplePay360 — HR Financial Export Engine
Generates Excel (.xlsx) and CSV exports for Payrun Payroll Registers, Bank ACH Transfer files, and Payslips.
"""
import io
import csv
from typing import List
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from models import Payrun, Payslip


def generate_payrun_register_df(payrun: Payrun, payslips: List[Payslip]) -> pd.DataFrame:
    """Builds a comprehensive DataFrame of payroll register lines per spec Section 5.1."""
    records = []
    for slip in payslips:
        emp = slip.employee
        lines_dict = {line.rule_code: line.amount for line in slip.lines}

        basic = lines_dict.get("BASIC", 0.0)
        hra = lines_dict.get("HRA", 0.0)
        conveyance = lines_dict.get("CONVEYANCE", 0.0)
        gross_before_lop = lines_dict.get("GROSS_BEFORE_LOP", basic + hra + conveyance)
        lop = lines_dict.get("LOP_DEDUCTION", 0.0)
        gross_after_lop = lines_dict.get("GROSS_AFTER_LOP", slip.gross_pay)
        pf = lines_dict.get("PF_DEDUCTION", 0.0)
        tax = lines_dict.get("TAX_DEDUCTION", 0.0)
        total_deduct = lines_dict.get("TOTAL_STATUTORY_DEDUCTIONS", slip.total_deductions)
        net_pay = lines_dict.get("NET_SALARY", slip.net_pay)

        contract_wage = slip.contract.wage if slip.contract else 0.0

        records.append({
            "Payrun Reference": payrun.reference,
            "Employee ID": emp.badge_id if emp else "",
            "Employee Name": f"{emp.first_name} {emp.last_name}" if emp else "",
            "Department": emp.department.name if emp and emp.department else "",
            "Job Position": emp.job_position.title if emp and emp.job_position else "",
            "Contract Base Wage ($)": contract_wage,
            "Scheduled Days": slip.scheduled_days,
            "Worked Days": slip.worked_days,
            "Unpaid Leave Days (LOP)": slip.unpaid_leave_days,
            "Basic Earnings ($)": basic,
            "House Rent Allowance ($)": hra,
            "Conveyance Allowance ($)": conveyance,
            "Gross Salary Before LOP ($)": gross_before_lop,
            "Loss of Pay Deduction ($)": lop,
            "Gross Salary After LOP ($)": gross_after_lop,
            "Provident Fund (PF) ($)": pf,
            "Income Tax / TDS ($)": tax,
            "Total Statutory Deductions ($)": total_deduct,
            "Net Salary Payable ($)": net_pay,
            "Disbursement Bank Name": emp.bank_name if emp else "",
            "Bank Account Number": emp.bank_account_no if emp else "",
            "IFSC / SWIFT Code": emp.ifsc_swift if emp else "",
            "Payout Status": slip.status.value if hasattr(slip.status, "value") else str(slip.status),
            "Verification Status": slip.verification_status.value if hasattr(slip.verification_status, "value") else str(slip.verification_status),
        })

    columns = [
        "Payrun Reference", "Employee ID", "Employee Name", "Department",
        "Job Position", "Contract Base Wage ($)", "Scheduled Days", "Worked Days",
        "Unpaid Leave Days (LOP)", "Basic Earnings ($)", "House Rent Allowance ($)",
        "Conveyance Allowance ($)", "Gross Salary Before LOP ($)", "Loss of Pay Deduction ($)",
        "Gross Salary After LOP ($)", "Provident Fund (PF) ($)", "Income Tax / TDS ($)",
        "Total Statutory Deductions ($)", "Net Salary Payable ($)", "Disbursement Bank Name",
        "Bank Account Number", "IFSC / SWIFT Code", "Payout Status", "Verification Status"
    ]
    return pd.DataFrame(records, columns=columns if not records else None)


def export_payrun_excel(payrun: Payrun, payslips: List[Payslip]) -> io.BytesIO:
    """Exports a formatted Excel workbook with styling."""
    df = generate_payrun_register_df(payrun, payslips)
    output = io.BytesIO()

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Payroll Register")
        worksheet = writer.sheets["Payroll Register"]

        # Styling
        header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        border = Border(
            left=Side(style='thin', color='D1D5DB'),
            right=Side(style='thin', color='D1D5DB'),
            top=Side(style='thin', color='D1D5DB'),
            bottom=Side(style='thin', color='D1D5DB')
        )

        for col_idx, col in enumerate(worksheet.iter_cols(min_row=1, max_row=1), start=1):
            for cell in col:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for row in worksheet.iter_rows(min_row=2):
            for cell in row:
                cell.border = border
                if isinstance(cell.value, float):
                    cell.number_format = '#,##0.00'

        for col in worksheet.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            worksheet.column_dimensions[col_letter].width = max(max_len + 4, 12)

    output.seek(0)
    return output


def export_payrun_csv(payrun: Payrun, payslips: List[Payslip]) -> str:
    """Exports CSV text representation of payroll register."""
    df = generate_payrun_register_df(payrun, payslips)
    return df.to_csv(index=False)


def export_bank_ach_csv(payrun: Payrun, payslips: List[Payslip]) -> str:
    """Exports Bank ACH / Direct NEFT disbursement file per spec Section 5.2."""
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow([
        "Beneficiary Name", "Bank Name", "Account Number", "IFSC/SWIFT",
        "Net Amount", "Payment Reference", "Remarks"
    ])

    for slip in payslips:
        emp = slip.employee
        name = f"{emp.first_name} {emp.last_name}" if emp else "Unknown"
        bank = emp.bank_name if emp and emp.bank_name else "N/A"
        acct = emp.bank_account_no if emp and emp.bank_account_no else "N/A"
        ifsc = emp.ifsc_swift if emp and emp.ifsc_swift else "N/A"
        ref = f"SAL-{payrun.period_start.strftime('%Y-%m')}-{emp.first_name.upper()}" if emp else "SAL-PAY"
        remarks = f"{payrun.name} Salary"

        writer.writerow([name, bank, acct, ifsc, f"{slip.net_pay:.2f}", ref, remarks])

    return output.getvalue()
